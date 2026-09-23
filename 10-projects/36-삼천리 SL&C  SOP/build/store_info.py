# -*- coding: utf-8 -*-
"""매장 정보.pptx → src/data.json 의 STORE FACT SHEET 페이지 채우기.

원본 입문매뉴얼의 'STORE FACT SHEET'는 빈칸 서식이라 data.json 에도 질문(head)만
있고 답(lines)이 비어 있다. 매장에서 작성한 '매장 정보.pptx'(1장=KSC, 2장=WASA)를
읽어 답을 채우고, 답이 'X'(해당 없음/미정)인 질문은 **질문째로 삭제**한다.

원본 PPTX 가 Fasoo DRM 으로 재암호화되면 이 스크립트는 못 읽는다.
그 때는 아래 FALLBACK 사전을 손으로 고쳐 쓰면 된다(--fallback).

사용법:
    python build/store_info.py            # 매장 정보.pptx 읽어 반영
    python build/store_info.py --fallback # PPTX 없이 아래 사전으로 반영
"""
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, 'src', 'data.json')
PPTX = os.path.join(
    os.path.expanduser('~'), 'Desktop', '교육팀',
    '4. 신규매장매뉴얼, 교안', '오픈매장 매뉴얼', 'KSC_US', '교육 자료', '매장 정보.pptx')

SLIDE_BRAND = ['KSC', 'WASA']     # 슬라이드 순서 = 브랜드 순서
PAGE_TITLE = 'STORE FACT SHEET'

# 주방 마감 시각 — 응대 멘트의 ____ 를 채운다 (PPTX 의 '라스트콜' 답에서 자동 추출)
CLOSE_FALLBACK = {'KSC': '(일~목) 21:00 (금,토) 22:00',
                  'WASA': '(일~목) 21:30 (금,토) 22:30'}

# PPTX 를 못 읽을 때 쓰는 수동 사전 (2026-09-23 기준)
FALLBACK = {
    'KSC': {
        'Store Name': 'Kalbi Social Club',
        'Address': '1510 Brea Mall Building A#1510, Brea, CA, 92821',
        'Phone': '714-786-6386',
        'Hours — Sun–Thu': '11:00~22:00',
        'Hours — Fri–Sat': '11:00~23:00',
        'Last Call & Kitchen Close': '(일~목) 21:00 (금,토) 22:00',
        'General Manager': 'Hye Bin Kim',
        'Assistant GM': 'Benjamin Cha',
        'Kitchen Manager': 'Hyungsun Do',
        'Seat Count & Sections': '138 seats',
        'Private / Large Party Capacity': '3 private rooms',
        'Target Average Check': '$65',
        'Reservation Platform': 'Opentable',
        'POS System': 'SpotOn',
        'Waitlist System': 'Opentable',
        'Wi-Fi Name & Password': 'KSC Guest / happy123',
    },
    'WASA': {
        'Store Name': 'Izakaya Wasa',
        'Address': '1510 Brea Mall Building A#1510, Brea, CA, 92821',
        'Phone': '714-882-5409',
        'Hours — Sun–Thu': '11:00~22:00',
        'Hours — Fri–Sat': '11:00~23:00',
        'Last Call & Kitchen Close': '(일~목) 21:30 (금,토) 22:30',
        'Assistant GM': 'Jongho Lim',
        'Kitchen Manager': 'James Seitz',
        'Seat Count & Sections': '44 seats',
        'Target Average Check': '$40',
        'Reservation Platform': 'Opentable',
        'POS System': 'SpotOn',
        'Waitlist System': 'Opentable',
        'Wi-Fi Name & Password': 'Wasa Guest / happy123',
    },
}


def key_of(head):
    """'Store Name / 매장명' → 'Store Name'. 라벨 표기가 조금 달라도 맞물리게 앞부분만 쓴다."""
    k = head.split('/')[0].strip()
    # 'Private / Large Party Capacity / 단체석' 처럼 슬래시가 영문 안에 있는 경우
    if k in ('Private',):
        k = ' / '.join(x.strip() for x in head.split('/')[:2])
    return re.sub(r'\s+', ' ', k)


def read_pptx():
    """슬라이드마다 라벨 도형 바로 아래 붙은 값 도형을 짝지어 읽는다."""
    from pptx import Presentation
    pr = Presentation(PPTX)
    out = {}
    for idx, slide in enumerate(pr.slides):
        if idx >= len(SLIDE_BRAND):
            break
        shapes = []
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if t:
                shapes.append((round(sh.left / 914400, 2), round(sh.top / 914400, 2), t))
        vals = {}
        for i, (x, y, t) in enumerate(shapes):
            if '/' not in t or len(t) > 60 or '\n' in t:
                continue
            # 바로 아래(0.4인치 이내) 같은 열에 있는 도형이 그 라벨의 답이다
            for x2, y2, t2 in shapes[i + 1:]:
                if abs(x2 - x) < 0.05 and 0 < y2 - y < 0.4:
                    vals[key_of(t)] = t2.strip()
                    break
        out[SLIDE_BRAND[idx]] = vals
    return out


def close_time(ans):
    """'(일~목) 21:00 (금,토) 22:00' → (영문 멘트용, 국문 멘트용).

    영문 문장 안에 한글을 섞으면 화면에서 영/국 두 줄로 쪼개지므로
    (app.js 의 pairLines 가 첫 한글에서 줄을 나눈다) 영문은 영문으로만 쓴다."""
    hm = re.findall(r'(\d{1,2}):(\d{2})', ans)
    if len(hm) >= 2:
        def ampm(h, m):
            h, m = int(h), int(m)
            return '%d%s %s' % (h % 12 or 12, (':%02d' % m) if m else '', 'AM' if h < 12 else 'PM')
        en = '%s Sun–Thu / %s Fri–Sat' % (ampm(*hm[0]), ampm(*hm[1]))
    else:
        en = re.sub(r'[^\x00-\x7F]+', '', ans).strip() or ans
    kr = re.sub(r'\s+', ' ', ans.replace('(금,토)', '· (금·토)')).strip()
    return en, kr


LEAD = ['Know these by heart.', '아래 정보는 외워 두세요.']


def fill(page, vals, brand):
    keep, dropped = [], []
    for b in page['blocks']:
        head = b.get('head', '')
        if b.get('kind') == 'script':
            b['lines'] = script_lines(b['lines'], vals, brand)
            keep.append(b)
            continue
        if not head:                       # '빈칸을 채우세요' 안내 → 이제 다 채워졌다
            b['lines'] = list(LEAD)
            keep.append(b)
            continue
        # 라벨에 이미 시간이 박혀 있으면 답과 겹치니 떼어 낸다
        head = re.sub(r'\s+\d{1,2}:\d{2}\s*[–\-~]\s*\d{1,2}:\d{2}\s*$', '', head)
        b['head'] = head
        k = key_of(head)
        norm = lambda s: re.sub(r'\s+', '', s).lower()
        if any(norm(head) == norm(v) for v in vals.values()):
            # 답 텍스트박스에 삭제된 질문(비상구·AED 등)이 붙어 추출된 경우 → 통째로 버린다
            continue
        if k not in vals:
            # '매장정보반영' 원본(2026-09-23~)은 답을 라벨 옆 별도 텍스트박스로 그려 두어
            # 추출하면 답이 줄 없는 블록으로 따로 나온다 → 버리고 라벨 블록에 답을 채운다
            if not b.get('lines'):
                continue
            keep.append(b)                 # 새 서식에 없는 질문은 그대로 둔다
            continue
        v = vals[k]
        if v.strip().upper() in ('X', '-', ''):
            dropped.append(head)
            continue
        b['lines'] = [v]
        keep.append(b)
    page['blocks'] = keep
    return dropped


# 멘트 줄은 브랜드마다 '영문 · 국문'이 한 줄이기도, 두 줄로 쪼개져 있기도 하다
DROP_SCRIPT = ('parking:', 'restroom:', '주차:', '화장실:')
CLOSE_SCRIPT = ('kitchen close:', '주방 마감:')


def script_lines(lines, vals, brand):
    """주차·화장실 멘트는 해당 정보가 삭제됐으니 같이 뺀다. 주방 마감 시각은 채운다."""
    en, kr = close_time(vals.get('Last Call & Kitchen Close') or CLOSE_FALLBACK[brand])
    out = []
    for ln in lines:
        low = ln.lower().lstrip('"“＂ ')
        if low.startswith(DROP_SCRIPT):
            continue
        if low.startswith(CLOSE_SCRIPT):
            ln = ln.replace('closes at ____', 'closes at ' + en)
            ln = re.sub(r'주방이\s*____시?에', '주방이 ' + kr + '에', ln)
        out.append(ln)
    return out


def main():
    use_fb = '--fallback' in sys.argv
    if not use_fb:
        try:
            vals = read_pptx()
        except Exception as e:                       # DRM 재암호화 등
            print('PPTX 읽기 실패(%s) → FALLBACK 사전 사용' % e)
            vals, use_fb = FALLBACK, True
    else:
        vals = FALLBACK

    decks = json.load(open(DATA, encoding='utf-8'))
    for deck in decks:
        if deck.get('kind') != 'manual':
            continue
        v = vals.get(deck.get('brand'))
        if not v:
            continue
        for p in deck.get('pages', []):
            if PAGE_TITLE not in (p.get('title') or '').upper():
                continue
            p['badge'] = 'MEMORIZE / 숙지 필수'      # 더 이상 빈칸 서식이 아니다
            dropped = fill(p, v, deck['brand'])
            print('%-12s p.%-3s 채움 %2d개 / 삭제 %d개' % (
                deck['id'], p['n'],
                sum(1 for b in p['blocks'] if b.get('head') and b.get('lines')), len(dropped)))
            for d in dropped:
                print('    - 삭제:', d)

    with open(DATA, 'w', encoding='utf-8') as f:
        json.dump(decks, f, ensure_ascii=False, separators=(",", ":"))
    print('written', DATA, '(fallback)' if use_fb else '(pptx)')


if __name__ == '__main__':
    main()
