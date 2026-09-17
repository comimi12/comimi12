# -*- coding: utf-8 -*-
"""PPTX(KSC/WASA 교육자료) -> 구조화 JSON + 웹용 이미지 추출.

원본 덱은 생성기로 만들어져 도형 배치가 일정하다. 그 규칙을 그대로 읽는다:
  - 매뉴얼 페이지: 상단 타이틀 + 본문 블록(첫 문단=볼드 헤딩, 나머지=항목)
  - 메뉴 페이지: 카드 밴드(넓은 컨테이너 사각형) 단위로 1메뉴 = 1카드
"""
import io, os, re, sys, json, hashlib
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
IMGDIR = os.path.join(ROOT, "src", "img")
SRC = r"C:\Users\owner\Desktop\교육팀\4. 신규매장매뉴얼, 교안\오픈매장 매뉴얼"

DECKS = [
    dict(id="ksc-manual", brand="KSC", kind="manual",
         title="입문 · 서비스 매뉴얼", sub="TRAINING & SERVICE MANUAL",
         path=os.path.join(SRC, "KSC_US", "교육 자료", "KSC_BREA_입문매뉴얼.pptx")),
    dict(id="ksc-menu", brand="KSC", kind="menu",
         title="메뉴 SOP · 54종", sub="MENU TRAINING GUIDE",
         path=os.path.join(SRC, "KSC_US", "교육 자료", "KSC_menu_sop_메뉴교육자료.pptx")),
    dict(id="wasa-manual", brand="WASA", kind="manual",
         title="입문 · 서비스 매뉴얼", sub="TRAINING & SERVICE MANUAL",
         path=os.path.join(SRC, "WASA", "WASA_BREA_Training_Manual_A4.pptx")),
    dict(id="wasa-menu", brand="WASA", kind="menu",
         title="메뉴 SOP · 90종", sub="MENU MANUAL",
         path=os.path.join(SRC, "WASA", "WASA_menu_manual_전체90종.pptx")),
]

_seen = {}


def _has_alpha(im):
    if im.mode == "P":
        return "transparency" in im.info
    try:
        return im.getchannel("A").getextrema()[0] < 255
    except Exception:
        return False


def save_image(shape):
    """그림을 web용(긴 변 1000px)으로 저장하고 상대경로 반환. 내용 해시로 중복 제거."""
    try:
        blob = shape.image.blob
    except Exception:
        return None
    h = hashlib.sha1(blob).hexdigest()[:16]
    if h in _seen:
        return _seen[h]
    try:
        im = Image.open(io.BytesIO(blob))
    except Exception:
        return None
    ext = "png" if im.mode in ("RGBA", "LA", "P") and _has_alpha(im) else "jpg"
    if im.width > 1000 or im.height > 1000:
        im.thumbnail((1000, 1000), Image.LANCZOS)
    name = h + "." + ext
    out = os.path.join(IMGDIR, name)
    if not os.path.exists(out):
        if ext == "jpg":
            im.convert("RGB").save(out, "JPEG", quality=82, optimize=True)
        else:
            im.save(out, "PNG", optimize=True)
    rel = "img/" + name
    _seen[h] = rel
    return rel


def flat(shapes, out=None):
    """그룹을 풀어 도형을 평면 리스트로."""
    if out is None:
        out = []
    for sh in shapes:
        if sh.shape_type == 6:
            flat(sh.shapes, out)
        else:
            out.append(sh)
    return out


def inch(v):
    return Emu(v).inches if v is not None else 0.0


def paras(sh):
    """문단 -> [(text, bold, size)] (빈 줄 제거)."""
    res = []
    if not sh.has_text_frame:
        return res
    for p in sh.text_frame.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        runs = [r for r in p.runs if r.text.strip()]
        bold = bool(runs) and all(r.font.bold for r in runs)
        size = next((r.font.size.pt for r in runs if r.font.size), 0)
        res.append((t, bold, size))
    return res


def maxsize(sh):
    return max([s for _, _, s in paras(sh)] or [0])


# -- 매뉴얼 페이지 --------------------------------------------------

BLOCK_KIND = [
    (re.compile(r"^\u26a0|NEVER GUESS|THIS PAGE IS THE LAW|NON-NEGOTIABLE"), "warn"),
    (re.compile(r"^\u2713|ALWAYS"), "do"),
    (re.compile(r"^\u2715|^\u2717|NEVER"), "dont"),
    (re.compile(r"^\U0001f4ac|SCRIPT|\uc2a4\ud06c\ub9bd\ud2b8|\uba58\ud2b8"), "script"),
]


def block_kind(heading):
    for rx, k in BLOCK_KIND:
        if rx.search(heading):
            return k
    return "normal"


CH_RX = re.compile(r"CH\.\s*(\d{2})\s+(.+?)\s*$")


def grid_table(shapes, wide=()):
    """격자로 배치된 텍스트박스 무리(알러지 매트릭스 등)를 표로 복원.

    표 도형이 아니라 셀마다 개별 텍스트박스로 그려진 페이지가 있다.
    top이 같은 것끼리 행으로 묶고, 각 셀을 left가 가장 가까운 열에 넣는다.
    원본에 값이 없는 칸은 도형 자체가 없으므로 빈칸으로 채워 열을 맞춘다.
    """
    groups = {}
    for s in shapes:
        groups.setdefault(round(inch(s.top) * 20), []).append(s)
    rows = [sorted(v, key=lambda s: inch(s.left))
            for _, v in sorted(groups.items()) if len(v) >= 3]
    if len(rows) < 5:
        return None, shapes

    # 열 기준선 = 셀이 가장 많은 행(머리글)의 left 좌표
    header = max(rows, key=len)
    cols = [inch(s.left) for s in header]

    def cell_text(s):
        return " ".join(t for t, _, _ in paras(s))

    out = []
    for r in rows:
        line = [""] * len(cols)
        for s in r:
            x = inch(s.left)
            j = min(range(len(cols)), key=lambda k: abs(cols[k] - x))
            line[j] = (line[j] + " " + cell_text(s)).strip()
        out.append(line)

    # 표 중간을 가르는 구분 띠(LUNCH / ALL-DAY 등)를 제자리에 끼워 넣는다
    top_of = {}
    for r, orig in zip(out, rows):
        top_of[id(r)] = inch(orig[0].top)
    seps = [(inch(s.top), [cell_text(s)]) for s in wide]
    merged = sorted([(top_of[id(r)], r) for r in out] + seps, key=lambda x: x[0])

    used = {id(s) for r in rows for s in r} | {id(s) for s in wide}
    return [r for _, r in merged], [s for s in shapes if id(s) not in used]


def parse_manual_slide(slide, n, W, H):
    shapes = flat(slide.shapes)
    pics, texts, tables = [], [], []
    for sh in shapes:
        st = str(sh.shape_type)
        if "PICTURE" in st:
            pics.append(sh)
        elif getattr(sh, "has_table", False):
            tables.append(sh)
        elif sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh)

    chapter = ""
    ch_no = ""

    def is_chrome(sh):
        """썸탭(우측 세로 번호), 푸터, 쪽번호 = 페이지 장식."""
        nonlocal chapter, ch_no
        L, T, Wd = inch(sh.left), inch(sh.top), inch(sh.width)
        if L > W - 0.55:
            return True
        if T > H - 0.85:
            m = CH_RX.search(paras(sh)[0][0]) if paras(sh) else None
            if m:
                ch_no, chapter = m.group(1), m.group(2)
            return True
        ps = paras(sh)
        if len(ps) == 1 and re.fullmatch(r"\d{1,3}", ps[0][0]) and Wd < 1.3:
            return True
        return False

    body = [s for s in texts if not is_chrome(s)]
    if not body:
        return None

    top_zone = [s for s in body if inch(s.top) < 1.15]
    title_sh = max(top_zone, key=maxsize) if top_zone else max(body, key=maxsize)
    title = paras(title_sh)[0][0]
    rest = [s for s in body if s is not title_sh]

    badge = ""
    ty = inch(title_sh.top)
    for s in list(rest):
        if abs(inch(s.top) - ty) < 0.25 and inch(s.left) > inch(title_sh.left) + 1.0:
            ps = paras(s)
            if ps and len(ps) == 1 and len(ps[0][0]) < 80:
                badge = " · ".join(x.strip() for x in ps[0][0].split("\n") if x.strip())
                rest.remove(s)
                break

    # 챕터 표지: 큰 번호 + 챕터명 + 설명만 있는 페이지
    is_divider = len(rest) <= 3 and re.fullmatch(r"\d{2}", title) is not None
    if is_divider:
        ordered = sorted(rest, key=lambda s: inch(s.top))
        lines = [paras(s)[0][0] for s in ordered if paras(s)]
        return dict(n=n, title=(lines[0] if lines else title), badge="CH." + title,
                    chapter=chapter or (lines[0] if lines else ""), ch=title or ch_no,
                    divider=True, blocks=[], images=[], tables=[], checklist=[],
                    sub=(lines[1] if len(lines) > 1 else ""),
                    intro=(lines[2] if len(lines) > 2 else ""))

    # 체크리스트 페이지: 원본이 CL_card / CL_bartitle / CL_item 로 이름 붙여 그렸다
    checklist = []
    if any(s.name.startswith("CL_bartitle") for s in rest):
        cards = sorted([s for s in shapes if s.name.startswith("CL_card")],
                       key=lambda s: (round(inch(s.left), 1), inch(s.top)))
        titles = [s for s in rest if s.name.startswith("CL_bartitle")]
        items = [s for s in rest if s.name.startswith("CL_item")]
        for c in cards:
            # 카드 제목 막대는 카드 상단보다 살짝 위에 걸쳐 있는 덱이 있어 위쪽을 넉넉히 잡는다
            x0, y0 = inch(c.left) - 0.05, inch(c.top) - 0.3
            x1, y1 = inch(c.left) + inch(c.width) + 0.05, inch(c.top) + inch(c.height) + 0.1

            def within(s):
                return x0 <= inch(s.left) <= x1 and y0 <= inch(s.top) <= y1

            ts = [s for s in titles if within(s)]
            its = sorted([s for s in items if within(s)], key=lambda s: inch(s.top))
            if not ts:
                continue
            checklist.append(dict(
                title=paras(ts[0])[0][0],
                items=[[t for t, _, _ in paras(s)] for s in its]))
        used = {id(s) for s in titles + items}
        rest = [s for s in rest if id(s) not in used]

    # 표: 진짜 표 도형 + 셀을 텍스트박스로 그린 격자
    tbls = []
    for t in tables:
        tbls.append([[c.text.strip() for c in r.cells] for r in t.table.rows])
    small = [s for s in rest if inch(s.height) < 0.45 and inch(s.width) < 3.0]
    if len(small) >= 30:
        ys = [inch(s.top) for s in small]
        # 표 구간을 가로지르는 넓은 띠 = 구분 행 (LUNCH / ALL-DAY 등)
        wide = [s for s in rest if inch(s.width) > 4.0 and inch(s.height) < 0.45
                and min(ys) <= inch(s.top) <= max(ys)]
        grid, leftover = grid_table(small, wide)
        if grid:
            tbls.append(grid)
            dropped = {id(x) for x in small} | {id(x) for x in wide}
            keep = {id(s) for s in leftover}
            rest = [s for s in rest if id(s) not in dropped or id(s) in keep]

    blocks = []
    for s in sorted(rest, key=lambda s: (round(inch(s.top), 1), round(inch(s.left), 1))):
        ps = paras(s)
        if not ps:
            continue
        L, T, Wd, Ht = inch(s.left), inch(s.top), inch(s.width), inch(s.height)
        meta = Ht < 0.62 and Wd < 2.6 and len(ps) == 2
        # 첫 문단이 제목인 조건: 볼드거나, 뒤 문단보다 크거나, 라벨+값 형태의 작은 칩.
        # 같은 크기의 비볼드 문단은 본문 도입부이므로 제목으로 올리지 않는다.
        head, lines = "", []
        if ps[0][1] or (len(ps) > 1 and ps[0][2] > ps[1][2]) or meta:
            head = ps[0][0]
            lines = [t for t, _, _ in ps[1:]]
        else:
            lines = [t for t, _, _ in ps]
        blocks.append(dict(head=head, lines=lines,
                           kind="meta" if meta else block_kind(head),
                           x=round(L, 2), y=round(T, 2), w=round(Wd, 2)))

    images = []
    for p in pics:
        if inch(p.top) > H - 0.85 or inch(p.left) > W - 0.55:
            continue
        rel = save_image(p)
        if rel:
            images.append(dict(src=rel, y=round(inch(p.top), 2), x=round(inch(p.left), 2),
                               w=round(inch(p.width), 2), h=round(inch(p.height), 2)))
    return dict(n=n, title=title, badge=badge, chapter=chapter, ch=ch_no, divider=False,
                blocks=blocks, images=images, tables=tbls, checklist=checklist)


# -- 메뉴 페이지 ----------------------------------------------------

SEC_KEYS = [
    ("SET INCLUDES", "구성"), ("COMPONENTS", "구성"),
    ("SERVING POINT", "제공 포인트"), ("DESCRIPTION", "메뉴 설명"),
    ("SUGGESTIVE SELLING", "추천 스크립트"), ("ALLERGENS", "알레르기"),
]


def sec_key(t):
    u = t.upper()
    for en, kr in SEC_KEYS:
        if u.startswith(en):
            return en, kr
    return None, None


def parse_menu_slide(slide, n, W, H):
    shapes = flat(slide.shapes)
    bands = [s for s in shapes
             if "PICTURE" not in str(s.shape_type)
             and inch(s.width) > W * 0.82 and inch(s.height) > 1.8]
    bands = sorted(bands, key=lambda s: inch(s.top))
    if not bands:
        return []
    cards = []
    for b in bands:
        t0 = inch(b.top) - 0.05
        t1 = inch(b.top) + inch(b.height) + 0.05
        # 카드 밴드 안쪽만. 페이지 오른쪽 세로 인덱스 탭은 밴드 밖이므로 중심점으로 걸러낸다.
        x0, x1 = inch(b.left) - 0.05, inch(b.left) + inch(b.width) + 0.05
        inside = [s for s in shapes
                  if s is not b and t0 <= inch(s.top) <= t1
                  and x0 <= inch(s.left) + inch(s.width) / 2 <= x1]
        pics = [s for s in inside if "PICTURE" in str(s.shape_type)]
        texts = [s for s in inside if s.has_text_frame and s.text_frame.text.strip()]
        if not texts:
            continue

        price, chip, code = "", "", ""
        body_sh = None
        for s in texts:
            ps = paras(s)
            if not ps:
                continue
            t = ps[0][0]
            if len(ps) == 1 and t.startswith("$"):
                price = t
            elif len(ps) == 1 and re.fullmatch(r"\d{1,3}", t) and inch(s.left) > W * 0.7:
                code = t
            elif len(ps) >= 4 and any(sec_key(p[0])[0] for p in ps):
                body_sh = s

        named = [s for s in texts
                 if s is not body_sh and len(paras(s)) == 1
                 and not paras(s)[0][0].startswith("$")
                 and paras(s)[0][0] != code]
        named = [s for s in named if "촬영 예정" not in paras(s)[0][0]]
        if not named:
            continue
        name_sh = max(named, key=maxsize)
        name_en = paras(name_sh)[0][0]
        nsz = maxsize(name_sh)

        chips = [s for s in named if s is not name_sh and inch(s.width) < 2.0
                 and inch(s.top) <= inch(name_sh.top) + 0.1 and maxsize(s) < nsz]
        if chips:
            chip = paras(min(chips, key=lambda s: inch(s.top)))[0][0]

        below = [s for s in named if s is not name_sh
                 and inch(s.top) > inch(name_sh.top)
                 and inch(s.left) >= inch(name_sh.left) - 0.2 and maxsize(s) < nsz]
        name_kr = paras(min(below, key=lambda s: inch(s.top)))[0][0] if below else ""

        sections, cur = [], None
        for t, bold, sz in (paras(body_sh) if body_sh else []):
            en, kr = sec_key(t)
            if en:
                cur = dict(en=en, kr=kr, lines=[])
                sections.append(cur)
            elif cur is not None:
                cur["lines"].append(t)

        if not sections:
            continue  # 요약/인덱스 표 페이지 — 메뉴 카드가 아니다
        img = None
        if pics:
            p = max(pics, key=lambda s: inch(s.width) * inch(s.height))
            img = save_image(p)
        cards.append(dict(page=n, code=code, chip=chip, en=name_en, kr=name_kr,
                          price=price, img=img, pending=img is None, sections=sections))
    return cards


def main():
    os.makedirs(IMGDIR, exist_ok=True)
    out = []
    for d in DECKS:
        if not os.path.exists(d["path"]):
            print("[!] 원본 없음: " + d["path"])
            continue
        prs = Presentation(d["path"])
        W, H = inch(prs.slide_width), inch(prs.slide_height)
        deck = {k: d[k] for k in ("id", "brand", "kind", "title", "sub")}
        deck["source"] = os.path.basename(d["path"])
        if d["kind"] == "menu":
            items = []
            for i, s in enumerate(prs.slides, 1):
                if i == 1:
                    continue
                items.extend(parse_menu_slide(s, i, W, H))
            deck["items"] = items
            print("  %-16s 메뉴 %d종 / %dp" % (d["id"], len(items), len(prs.slides)))
        else:
            pages = [p for p in (parse_manual_slide(s, i, W, H)
                                 for i, s in enumerate(prs.slides, 1)) if p]
            deck["pages"] = pages
            print("  %-16s 페이지 %d / %dp" % (d["id"], len(pages), len(prs.slides)))
        out.append(deck)
    dst = os.path.join(ROOT, "src", "data.json")
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    print("[OK] %s (%.0f KB) 이미지 %d개"
          % (dst, os.path.getsize(dst) / 1024, len(os.listdir(IMGDIR))))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
